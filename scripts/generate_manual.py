#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成系统使用手册 .docx 文档。"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

doc = Document()

style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.3

# ── 封面 ──
title = doc.add_heading("轻量化多模态轨迹预测系统 使用手册", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("面向低速场景 · 实体小车部署指南")
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
m = doc.add_paragraph()
m.alignment = WD_ALIGN_PARAGRAPH.CENTER
m.add_run(f"文档版本: 1.0 | 日期: {datetime.date.today().strftime('%Y-%m-%d')}").font.size = Pt(10)
doc.add_page_break()

# ═══ 一、系统概述 ═══
doc.add_heading("一、系统概述", level=1)
doc.add_heading("1.1 项目定位", level=2)
doc.add_paragraph(
    "本系统是一个面向低速场景（校园/园区无人车）的轻量化多模态轨迹预测系统。"
    "核心目标：用最少的参数（195K）和最低的算力（RTX 3060 6GB），"
    "在 Argoverse 真实数据上做出可用的多模态轨迹预测。"
    "从数据管道、模型训练、消融实验、可视化分析到 ONNX 部署全流程覆盖，"
    "可直接部署到 Intel NUC 实体小车。"
)

doc.add_heading("1.2 系统架构", level=2)
doc.add_paragraph(
    "数据流: Argoverse CSV -> preprocessing.py(11维特征) -> train.npz -> "
    "TrajectoryModel(K=5+时序卷积+注意力池化) -> WTA训练(AMP) -> "
    "ONNX导出(0.77MB) -> ROS节点(20Hz) -> RViz可视化"
)

doc.add_heading("1.3 核心技术指标", level=2)
t = doc.add_table(rows=8, cols=2, style="Light Grid Accent 1")
for i, (k, v) in enumerate([
    ("模型参数量", "194,733 (目标 200K 以内)"),
    ("模态数 K", "5"),
    ("输入特征", "11维 [x,y,vx,vy,t/T,1,dist,rel_x,rel_y,heading,curvature]"),
    ("观测/预测窗口", "20帧(2s) / 30帧(3s) @10Hz"),
    ("验证集 minADE", "0.5779 (80 epoch)"),
    ("CPU 推理延迟", "0.21 ms (ONNX Runtime)"),
    ("ONNX 模型大小", "0.77 MB"),
    ("训练硬件", "RTX 3060 6GB, ~4h (80 epoch)"),
]):
    t.rows[i].cells[0].text = k
    t.rows[i].cells[1].text = v

doc.add_heading("1.4 模型架构", level=2)
doc.add_paragraph(
    "编码器: Linear(11->128)+LayerNorm -> 3xMLPBlock(残差) -> "
    "DepthwiseConv1d(时序混合) -> AttentionPooling(标准1/sqrt(d_k))\n"
    "预测器: 模态选择器(128->64->5) + K=5个轨迹分支(128->128->60)\n"
    "不确定性: Linear(128->64->60)+Softplus -> 每帧每维方差\n"
    "损失: L = 0.7*WTA_smooth_l1 + 0.3*CE + 0.05*NLL(可选)"
)

doc.add_page_break()

# ═══ 二、训练 ═══
doc.add_heading("二、训练流程", level=1)
doc.add_heading("2.1 数据准备", level=2)
doc.add_paragraph("1. 将 Argoverse v1.1 解压到 1/train/data/ 和 1/val/data/")
doc.add_paragraph("2. python scripts/preprocess_real_data.py (处理约245K个CSV)")
doc.add_paragraph("3. 输出 data/processed/train.npz(150万样本) val.npz(21万) test.npz(43万)")

doc.add_heading("2.2 训练命令", level=2)
doc.add_paragraph("完整训练: python experiments/run_baseline.py (80 epoch, ~4h)")
doc.add_paragraph("快速消融: python scripts/run_ablation_continue.py (5 epoch x 8变体)")
doc.add_paragraph("两阶段不确定性: python scripts/train_with_uncertainty.py")

doc.add_heading("2.3 训练策略", level=2)
doc.add_paragraph(
    "优化器: AdamW lr=1e-3 | 调度: CosineAnnealing+5epoch warmup | "
    "混合精度: AMP | 增强: 旋转+缩放+噪声 | 早停: patience=15"
)

doc.add_page_break()

# ═══ 三、消融与提升 ═══
doc.add_heading("三、消融实验与性能提升", level=1)

doc.add_heading("3.1 消融结果", level=2)
t = doc.add_table(rows=5, cols=3, style="Light Grid Accent 1")
for i, h in enumerate(["变体", "minADE", "vs基线"]):
    t.rows[0].cells[i].text = h
for i, (a, b, c) in enumerate([
    ("K=3 基线", "0.7629", "--"),
    ("K=5 ★", "0.6334", "+17.0%"),
    ("无上下文6维", "0.7472", "+2.1%"),
    ("K=1", "1.1539", "-51.2%"),
]):
    t.rows[i+1].cells[0].text = a
    t.rows[i+1].cells[1].text = b
    t.rows[i+1].cells[2].text = c

doc.add_heading("3.2 四阶段提升路线", level=2)
t = doc.add_table(rows=5, cols=3, style="Light Grid Accent 1")
for i, h in enumerate(["阶段", "minADE", "累计提升"]):
    t.rows[0].cells[i].text = h
for i, (a, b, c) in enumerate([
    ("起点: K=3,9维,5ep", "0.7589", "--"),
    ("+11维特征", "0.7655", "-0.9%"),
    ("+K=5+temporal_conv", "0.6485", "+14.5%"),
    ("+80 epoch", "0.5779", "+23.9%"),
]):
    t.rows[i+1].cells[0].text = a
    t.rows[i+1].cells[1].text = b
    t.rows[i+1].cells[2].text = c

doc.add_heading("3.3 关键发现", level=2)
for f in [
    "K=5 是最大单一杠杆: +17% 提升，代价仅 48K 参数",
    "交互特征对轻量模型是噪声: 6维去上下文反而优于9维",
    "池化方式几乎无影响: 均值=注意力=LSTM (差值<0.01)",
    "延长训练有确定收益: 80 epoch 比 5 epoch 额外提升 10.9%",
]:
    doc.add_paragraph(f, style="List Bullet")

doc.add_page_break()

# ═══ 四、Web Demo ═══
doc.add_heading("四、Web 可视化 Demo", level=1)
doc.add_paragraph("支持手绘轨迹、上传 Argoverse CSV、内置示例三种输入方式。K=5 模态预测结果以彩色轨迹+不确定性椭圆展示。")

doc.add_heading("4.1 启动", level=2)
doc.add_paragraph("cd da_chuang && python demo/app.py\n浏览器打开 http://localhost:5000")

doc.add_heading("4.2 操作", level=2)
for op in [
    "左键点击: 添加轨迹点(至少20个) | 左键拖拽: 微调节点",
    "右键拖拽: 平移画布 | 滚轮: 缩放",
    "开始预测: 取前20个点推理，显示5条预测轨迹",
    "上传CSV: 上传Argoverse格式CSV，自动解析agent轨迹",
    "内置示例: 直线/弯道/掉头/换道 4个预设场景",
]:
    doc.add_paragraph(op, style="List Bullet")

doc.add_heading("4.3 API", level=2)
doc.add_paragraph(
    "POST /api/v1/predict — 坐标序列预测\n"
    "POST /api/v1/predict_csv — CSV上传预测\n"
    "GET /api/v1/examples — 内置示例\n"
    "GET /api/v1/model_info — 模型信息"
)

doc.add_page_break()

# ═══ 五、小车部署 ═══
doc.add_heading("五、实体小车部署", level=1)

doc.add_heading("5.1 小车硬件", level=2)
doc.add_paragraph(
    "上位机: Intel NUC (Ubuntu 20.04+ROS Noetic)\n"
    "底盘: Arduino Mega + 麦克纳姆轮(全向)\n"
    "传感器: RPLidar+深度摄像头+编码器+IMU\n"
    "导航速度: 0.015~0.12 m/s | 控制频率: 20 Hz\n"
    "WiFi: SSID rak_xxx 密码 12345678 | SSH: mo@10.42.0.1"
)

doc.add_heading("5.2 部署步骤", level=2)
for s in [
    "1. 拷贝ONNX模型: scp deploy/trajectory_model.onnx mo@10.42.0.1:~/ros_space/src/ai_node_master/models/",
    "2. 拷贝ROS节点: scp deploy/ros_node/trajectory_predictor.py mo@10.42.0.1:~/ros_space/src/ai_node_master/src/",
    "3. SSH安装依赖: ssh mo@10.42.0.1 && pip3 install onnxruntime numpy",
    "4. 启动导航: roslaunch ai_node_master navigation.launch",
    "5. 启动预测: cd ~/ros_space && source devel/setup.bash && python3 src/ai_node_master/src/trajectory_predictor.py",
    "6. RViz查看: Add -> By topic -> /predicted_trajectories -> MarkerArray",
]:
    doc.add_paragraph(s)

doc.add_heading("5.3 话题接口", level=2)
doc.add_paragraph(
    "/odom -> 订阅(20Hz) 里程计输入\n"
    "/predicted_trajectories -> 发布(10Hz) 预测路径可视化"
)

doc.add_heading("5.4 预测节点数据流", level=2)
doc.add_paragraph(
    "/odom(20Hz) -> 提取(x,y,yaw)存20帧 -> 每10Hz触发预测 "
    "-> 末帧居中归一化 -> 11维特征 -> ONNX推理(0.21ms) "
    "-> 5条轨迹还原绝对坐标 -> MarkerArray"
)

doc.add_page_break()

# ═══ 六、文件结构 ═══
doc.add_heading("六、项目文件结构", level=1)
doc.add_paragraph(
    "configs/          YAML配置\n"
    "data/             数据管道(预处理+Dataset+增强)\n"
    "models/           模型(编码器+池化+预测器+不确定性)\n"
    "train/            训练(损失+优化器+Trainer)\n"
    "eval/             评估(指标+评估器)\n"
    "experiments/      实验脚本\n"
    "visualization/    论文图表生成(fig3-10)\n"
    "scripts/          工具脚本\n"
    "demo/             Web可视化\n"
    "deploy/           ONNX+ROS部署包\n"
    "outputs/          输出(checkpoints+logs+figures+results)"
)

doc.add_page_break()

# ═══ 七、论文材料 ═══
doc.add_heading("七、论文支撑材料", level=1)
doc.add_paragraph(
    "图表 outputs/figures/: fig3(池化) fig4(训练曲线) fig5(Pareto) "
    "fig6(消融) fig7(延迟) fig8(轨迹) fig9(不确定性) fig10(失败案例)\n\n"
    "数据表 outputs/results/: table1(主结果) table2(场景分布) "
    "table3(消融) table4(效率) table5(参数扫描) table6(特征维度)\n\n"
    "GitHub: https://github.com/swee-theat/-"
)

output_path = "轻量化多模态轨迹预测系统_使用手册.docx"
doc.save(output_path)
print(f"文档已保存: {output_path}")
